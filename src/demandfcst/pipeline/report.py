"""Weekly data quality report (Excel) for the latest delivery.

It re-validates raw (so a delivery that failed processing is still reported), adds sanity
checks against history from the processed weekly table, lists deliveries and corrections,
shows weekly volumes by region with a chart, and records the code version and unit tests.
Status: FAIL (errors, failed tests, or processing missed the latest delivery),
PASS WITH WARNINGS, or PASS.
"""

import json
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from demandfcst import drops, storage
from demandfcst.pipeline import dedupe, reader, sanity, validate

FILLS = {"pass": "C6EFCE", "passed": "C6EFCE", "PASS": "C6EFCE",
         "info": "DDEBF7",
         "warning": "FFEB9C", "skipped": "FFEB9C", "PASS WITH WARNINGS": "FFEB9C",
         "error": "FFC7CE", "failure": "FFC7CE", "FAIL": "FFC7CE"}
LEVEL_ORDER = {"error": 0, "warning": 1, "info": 2, "pass": 3}
WEEKS_SHOWN = 12


def git_version(project_root):
    """Short commit hash of the code, flagged if there are uncommitted changes."""
    def git(*args):
        return subprocess.run(["git", *args], cwd=project_root, capture_output=True,
                              text=True, check=True).stdout.strip()
    try:
        commit = git("rev-parse", "--short", "HEAD")
        return commit + (" + uncommitted changes" if git("status", "--porcelain") else "")
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def unit_tests(xml_path):
    """(summary, cases) from a `pytest --junitxml` file, or (None, empty table) if there is none."""
    if not xml_path or not Path(xml_path).exists():
        return None, pd.DataFrame(columns=["test", "outcome"])
    root = ET.parse(xml_path).getroot()
    suites = [root] if root.tag == "testsuite" else root.findall("testsuite")
    summary = {k: sum(int(s.get(k, 0)) for s in suites) for k in ("tests", "failures", "errors", "skipped")}
    summary["run_at"] = suites[0].get("timestamp", "") if suites else ""
    cases = []
    for case in root.iter("testcase"):
        outcome = next((tag for tag in ("failure", "error", "skipped") if case.find(tag) is not None), "passed")
        cases.append({"test": f"{case.get('classname')}::{case.get('name')}", "outcome": outcome})
    return summary, pd.DataFrame(cases, columns=["test", "outcome"])


def _processed(processed_root, name):
    location = storage.join(processed_root, f"{name}.parquet")
    return pd.read_parquet(location) if storage.exists(location) else None


def build_report(raw_root, processed_root, tests_xml=None, project_root=None):
    """All report sheets for the latest delivery, plus its status."""
    tables = reader.read_all(raw_root)
    checks = validate.validate(tables)
    checks = checks.sort_values("level", key=lambda s: s.map(LEVEL_ORDER), kind="stable").reset_index(drop=True)
    latest = drops.drop_dates(raw_root, "sellin_daily")[-1]

    info_path = storage.join(processed_root, "run_info.json")
    info = json.loads(storage.read_text(info_path)) if storage.exists(info_path) else {}
    up_to_date = info.get("latest_drops", {}).get("sellin_daily") == f"{latest:%Y-%m-%d}"

    weekly = _processed(processed_root, "sales_weekly")
    restatements = _processed(processed_root, "restatements")
    if restatements is None:
        restatements = pd.DataFrame(columns=dedupe.SUMMARY_COLUMNS)
    sanity_df = (sanity.sanity_checks(weekly, tables["province_master"]) if weekly is not None
                 else pd.DataFrame(columns=sanity.COLUMNS))
    tests, test_cases = unit_tests(tests_xml)

    n = checks["level"].value_counts()
    errors = int(n.get("error", 0))
    warnings = int(n.get("warning", 0)) + int((sanity_df["status"] == "warning").sum())
    tests_failed = tests is not None and tests["failures"] + tests["errors"] > 0
    status = ("FAIL" if errors or tests_failed or not up_to_date
              else "PASS WITH WARNINGS" if warnings else "PASS")

    week_lines, volumes = [], pd.DataFrame()
    if weekly is not None:
        totals = weekly.groupby("week_start")["volume_hl"].sum()
        last = totals.index.max()
        week_lines = [("Latest complete week", f"{last:%Y-%m-%d} .. {last + pd.Timedelta(days=6):%Y-%m-%d}"),
                      ("Volume in that week", f"{totals[last]:,.0f} HL"),
                      ("Weekly table", f"{info.get('first_week')} .. {info.get('last_week')}: "
                                       f"{info.get('weeks')} weeks, {info.get('series')} series")]
        by_region = (weekly.merge(tables["province_master"][["province", "region"]], on="province")
                     .pivot_table(index="week_start", columns="region", values="volume_hl", aggfunc="sum"))
        by_region["Total"] = by_region.sum(axis=1)
        volumes = by_region.tail(WEEKS_SHOWN).round(1).reset_index()

    this_week = restatements[restatements["ingest_date"] == latest]
    fixes = ("none" if this_week.empty else
             f"{int(this_week['restated_days'].sum())} earlier days restated, {this_week['hl_change'].sum():+,.1f} HL")
    summary = pd.DataFrame([
        ("Delivery", f"week {latest - pd.Timedelta(days=6):%Y-%m-%d} .. {latest:%Y-%m-%d}"),
        ("Overall status", status),
        ("Validation checks", f"{len(checks)} run: {int(n.get('pass', 0))} passed, {int(n.get('info', 0))} info, "
                              f"{int(n.get('warning', 0))} warnings, {errors} errors"),
        ("Sanity checks", f"{len(sanity_df)} run: {int((sanity_df['status'] == 'pass').sum())} passed, "
                          f"{int((sanity_df['status'] == 'warning').sum())} warnings"),
        ("Processed data", f"up to date (processed {info.get('processed_at')})" if up_to_date
                           else "NOT processed for this delivery: run process.py"),
        *week_lines,
        ("Corrections in this delivery", fixes),
        ("Code version", git_version(project_root) if project_root else "unknown"),
        ("Unit tests", "not run" if tests is None else
                       f"{tests['tests']} run, {tests['failures'] + tests['errors']} failed, "
                       f"{tests['skipped']} skipped ({tests['run_at']})"),
        ("Generated at", pd.Timestamp.now().strftime("%Y-%m-%d %H:%M")),
    ], columns=["item", "value"])

    deliveries = (tables["sellin_daily"].groupby("ingest_date")
                  .agg(rows=("volume_hl", "size"), first_day=("date", "min"), last_day=("date", "max"),
                       total_hl=("volume_hl", "sum"))
                  .reset_index())
    deliveries["total_hl"] = deliveries["total_hl"].round(1)

    sheets = {"Summary": summary, "Checks": checks, "Sanity": sanity_df, "Deliveries": deliveries,
              "Corrections": restatements, "Weekly volumes": volumes, "Unit tests": test_cases}
    return {"latest": latest, "status": status, "sheets": sheets}


def _dates_as_text(df):
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = out[col].dt.strftime("%Y-%m-%d")
    return out


def _style(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            fill = FILLS.get(str(cell.value))
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(longest + 2, 90)
    ws.freeze_panes = "A2"


def _add_chart(ws):
    if ws.max_row < 3:
        return
    chart = LineChart()
    chart.title = "Weekly sell-in by region (HL)"
    chart.height, chart.width = 9, 22
    chart.add_data(Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=ws.max_row),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(chart, f"{get_column_letter(ws.max_column + 2)}2")


def write_excel(sheets, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            _dates_as_text(df).to_excel(xw, sheet_name=name, index=False)
        for ws in xw.book.worksheets:
            _style(ws)
        _add_chart(xw.book["Weekly volumes"])


def make_report(raw_root, processed_root, out_dir, tests_xml=None, project_root=None):
    """Build and write data_quality_<ingest_date>.xlsx. Returns (path, status)."""
    rep = build_report(raw_root, processed_root, tests_xml, project_root)
    path = storage.join(out_dir, f"data_quality_{rep['latest']:%Y-%m-%d}.xlsx")
    write_excel(rep["sheets"], path)
    return path, rep["status"]
